from docx import Document

from app.core.loaders import load_document


def test_docx_tables_are_preserved_as_markdown(tmp_path):
    path = tmp_path / "expense.docx"
    document = Document()
    document.add_heading("报销标准", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "住宿上限"
    table.cell(1, 0).text = "上海"
    table.cell(1, 1).text = "500元"
    document.save(path)

    pages = load_document(path)
    content = "\n".join(page.content for page in pages)

    assert "| 城市 | 住宿上限 |" in content
    assert "| 上海 | 500元 |" in content
