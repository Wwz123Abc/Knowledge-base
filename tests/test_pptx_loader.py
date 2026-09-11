from pptx import Presentation

from app.core.loaders import load_document


def test_pptx_text_and_table_are_indexable(tmp_path):
    path = tmp_path / "policy.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "差旅标准"
    textbox = slide.shapes.add_textbox(0, 0, 3000000, 500000)
    textbox.text = "住宿费用不得超过标准"
    table = slide.shapes.add_table(2, 2, 0, 600000, 3000000, 1000000).table
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "上限"
    table.cell(1, 0).text = "上海"
    table.cell(1, 1).text = "500元"
    presentation.save(path)

    pages = load_document(path)
    content = "\n".join(page.content for page in pages)

    assert "差旅标准" in content
    assert "住宿费用" in content
    assert "| 上海 | 500元 |" in content
